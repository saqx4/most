import csv

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def export_to_csv(queryset, fields, filename='export.csv'):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    headers = [f.replace('_', ' ').title() for f in fields]
    writer.writerow(headers)
    for obj in queryset:
        row = []
        for f in fields:
            value = obj
            for attr in f.split('__'):
                if value is None:
                    break
                value = getattr(value, attr, None)
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m-%d')
            elif hasattr(value, 'pk'):
                value = str(value)
            elif value is None:
                value = ''
            row.append(value)
        writer.writerow(row)
    return response


def export_to_excel(queryset, fields, filename='export.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Export'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')

    headers = [f.replace('_', ' ').title() for f in fields]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, obj in enumerate(queryset, 2):
        for col_idx, f in enumerate(fields, 1):
            value = obj
            for attr in f.split('__'):
                if value is None:
                    break
                value = getattr(value, attr, None)
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m-%d')
            elif hasattr(value, 'pk'):
                value = str(value)
            elif value is None:
                value = ''
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def import_customers_from_file(file_obj, company, user=None):
    """Import customers from CSV or Excel file."""
    from apps.sales.models import Customer
    from decimal import Decimal
    import csv, io, openpyxl

    created_count = 0
    errors = []

    filename = getattr(file_obj, 'name', '')
    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        wb = openpyxl.load_workbook(file_obj)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        data_rows = rows[1:]
    else:
        content = file_obj.read().decode('utf-8-sig', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() for h in rows[0]]
        data_rows = rows[1:]

    def get_val(row, *col_names):
        for name in col_names:
            if name in headers:
                idx = headers.index(name)
                if idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
        return ''

    for r_idx, row in enumerate(data_rows, start=2):
        if not any(row):
            continue
        name = get_val(row, 'name', 'اسم العميل', 'الاسم', 'customer_name')
        if not name:
            errors.append(f'صف {r_idx}: اسم العميل مفقود')
            continue

        email = get_val(row, 'email', 'البريد الالكتروني', 'البريد')
        phone = get_val(row, 'phone', 'الهاتف', 'رقم الهاتف')
        mobile = get_val(row, 'mobile', 'جوال', 'الموبايل')
        tax_id = get_val(row, 'tax_id', 'الرقم الضريبي', 'الضريبي')
        comm_reg = get_val(row, 'commercial_reg', 'السجل التجاري')
        client_type = get_val(row, 'client_type', 'نوع العميل') or 'commercial'

        cust, created = Customer.objects.get_or_create(
            company=company,
            name=name,
            defaults={
                'email': email,
                'phone': phone,
                'mobile': mobile,
                'tax_id': tax_id,
                'commercial_reg': comm_reg,
                'client_type': 'individual' if 'فرد' in client_type or 'ind' in client_type else 'commercial',
                'created_by': user
            }
        )
        if created:
            created_count += 1

    return created_count, errors


def import_products_from_file(file_obj, company, user=None):
    """Import products and inventory items from CSV or Excel file."""
    from apps.inventory.models import Product, Category, UnitOfMeasure
    from decimal import Decimal
    import csv, io, openpyxl

    created_count = 0
    errors = []

    filename = getattr(file_obj, 'name', '')
    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        wb = openpyxl.load_workbook(file_obj)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        data_rows = rows[1:]
    else:
        content = file_obj.read().decode('utf-8-sig', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() for h in rows[0]]
        data_rows = rows[1:]

    def get_val(row, *col_names):
        for name in col_names:
            if name in headers:
                idx = headers.index(name)
                if idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
        return ''

    for r_idx, row in enumerate(data_rows, start=2):
        if not any(row):
            continue
        sku = get_val(row, 'sku', 'كود الصنف', 'الرقم التسلسلي', 'كود')
        name = get_val(row, 'name', 'اسم المنتج', 'اسم الصنف', 'الاسم')
        if not name:
            errors.append(f'صف {r_idx}: اسم المنتج مفقود')
            continue
        if not sku:
            sku = f'SKU-{r_idx:04d}'

        sale_price_str = get_val(row, 'sale_price', 'سعر البيع', 'السعر') or '0'
        purch_price_str = get_val(row, 'purchase_price', 'سعر الشراء', 'التكلفة') or '0'
        barcode = get_val(row, 'barcode', 'الباركود')
        brand = get_val(row, 'brand', 'الماركة', 'العلامة التجارية')

        try:
            sale_price = Decimal(sale_price_str)
        except Exception:
            sale_price = Decimal('0.00')

        try:
            purch_price = Decimal(purch_price_str)
        except Exception:
            purch_price = Decimal('0.00')

        prod, created = Product.objects.get_or_create(
            company=company,
            sku=sku,
            defaults={
                'name': name,
                'sale_price': sale_price,
                'purchase_price': purch_price,
                'barcode': barcode,
                'brand': brand,
                'created_by': user
            }
        )
        if created:
            created_count += 1

    return created_count, errors


def import_suppliers_from_file(file_obj, company, user=None):
    """Import suppliers from CSV or Excel file."""
    from apps.purchasing.models import Supplier
    from decimal import Decimal
    import csv, io, openpyxl

    created_count = 0
    errors = []

    filename = getattr(file_obj, 'name', '')
    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        wb = openpyxl.load_workbook(file_obj)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() if h else '' for h in rows[0]]
        data_rows = rows[1:]
    else:
        content = file_obj.read().decode('utf-8-sig', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return 0, ['الملف فارغ']
        headers = [str(h).strip().lower() for h in rows[0]]
        data_rows = rows[1:]

    def get_val(row, *col_names):
        for name in col_names:
            if name in headers:
                idx = headers.index(name)
                if idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
        return ''

    for r_idx, row in enumerate(data_rows, start=2):
        if not any(row):
            continue
        name = get_val(row, 'name', 'اسم المورد', 'الاسم', 'supplier_name')
        if not name:
            errors.append(f'صف {r_idx}: اسم المورد مفقود')
            continue

        contact_person = get_val(row, 'contact_person', 'جهة الاتصال', 'الاتصال')
        email = get_val(row, 'email', 'البريد الالكتروني', 'البريد')
        phone = get_val(row, 'phone', 'الهاتف', 'رقم الهاتف')
        tax_id = get_val(row, 'tax_id', 'الرقم الضريبي', 'الضريبي')

        supplier, created = Supplier.objects.get_or_create(
            company=company,
            name=name,
            defaults={
                'contact_person': contact_person,
                'email': email,
                'phone': phone,
                'tax_id': tax_id,
                'created_by': user
            }
        )
        if created:
            created_count += 1

    return created_count, errors
